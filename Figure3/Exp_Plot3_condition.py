"""Recompute all four Plot_3 figures for Exp and Fine data.

The original ``Plot_3.ipynb`` plots condition numbers that were calculated
upstream and then copied into a dictionary/Excel table.  A condition number
cannot be corrected by scaling that final scalar.  This script therefore goes
back to the matrices on which the condition numbers are defined.

Two column-wise transformations are available for the reported condition
numbers.  The default is L2 normalization,

    Theta_scaled[:, j] = Theta[:, j] / ||Theta[:, j]||_2.

Pass ``--scaling std`` to use

    Theta_scaled[:, j] = Theta[:, j] / sample_std(Theta[:, j]),

where ``sample_std`` uses ``ddof=1``.  Thus every nonconstant column has
standard deviation one, while its mean is retained.  Pass ``--center`` to use
the centered z-score variant ``(Theta-mean)/sample_std`` instead.

Sparse regression is performed on the original, unscaled feature library.
After recovery identifies the wrong-or-missing term names, those columns are
evaluated on the original state data and only that fixed sublibrary is scaled
for its condition-number calculation.  Consequently, changing ``--scaling``
does not change the recovered support; it changes only the reported condition
numbers.

For each dataset, the four outputs reproduce the views from Plot_3:

1. full polynomial-library condition number versus degree;
2. full polynomial-library condition number versus model complexity;
3. wrong-or-missing-term sublibrary condition number versus degree; and
4. wrong-or-missing-term sublibrary condition number versus model complexity.

Exp and Fine use their own state data and recovered term sets, with the same
model-specific preprocessing and recovery settings. Hass has no separate
Exp/Fine files: both runs explicitly use its shared Data_used/Hass.xlsx file.
Excel detail sheets record the dataset, actual source file and whether it is shared.

Figures 3/4 must first repeat the model-recovery step because the old
``Table_fine.xlsx`` contains only already-computed condition-number scalars,
not the term subsets or their matrices.  This makes a complete run relatively
expensive, especially for the 15-state Zheng degree-5 library.  An optional
JSON term cache can skip recovery on later runs.

By default both datasets are processed and the figures are displayed. Use
``--save`` to write PNG/SVG figures and an Excel workbook separately into
``exp/`` and ``fine/`` output directories, plus a dataset-specific term cache.
Each workbook includes model-by-degree tables of condition numbers and their
log10 values, full-library details, and wrong/missing-term results, including
equation-wise mean/min/max on both the original condition and log10 scales.
``--full-only`` skips recovery and Figures 3/4.

Example: python Exp_Plot3_condition.py --save --no-show
"""

from __future__ import annotations

import argparse
import contextlib
import gc
import io
import json
import math
import os
import re
import sys
import tempfile
import time
import warnings
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pandas as pd


DEGREES = (2, 3, 4, 5)
DATASETS = ("exp", "fine")
DEGREE_COLORS = ("#0173B2", "#DE8F05", "#029E73", "#CC78BC")
DEGREE_MARKERS = ("o", "s", "^", "v")
DEGREE_LINESTYLES = ("-", "--", "-.", ":")

SCRIPT_DIR = Path(__file__).resolve().parent
REPOSITORY_ROOT = SCRIPT_DIR.parent
DEFAULT_SINDY_ROOT = REPOSITORY_ROOT.parent
DEFAULT_OUTPUT_DIR = REPOSITORY_ROOT / "results" / "plot3_condition_scaling"


@dataclass(frozen=True)
class ModelConfig:
    name: str
    plotted_complexity: int
    data_path: str
    model_path: str
    smooth_window: int
    smooth_order: int
    full_library_uses_smoothed_states: bool
    differential_order: int
    thresholds: tuple[float, float, float, float]
    alpha: float | None = None
    after_threshold: float | None = None
    recover_from_unsmoothed_data: bool = False
    drop_state_columns: tuple[str, ...] = ()
    minimum_time: float | None = None
    dataset: str = "fine"
    data_source: str = "fine"

    def threshold(self, degree: int) -> float:
        return self.thresholds[DEGREES.index(degree)]


MODEL_CONFIGS = (
    ModelConfig(
        name="Beer",
        plotted_complexity=4,
        data_path="Data/Data_fine/Beer_fine.xlsx",
        model_path="Model_Beer.xlsx",
        smooth_window=11,
        smooth_order=3,
        full_library_uses_smoothed_states=False,
        differential_order=2,
        thresholds=(1e-7, 1e-7, 1e-7, 1e-7),
        after_threshold=1e-6,
        minimum_time=1126.30654465543 - 5.0,
    ),
    ModelConfig(
        name="Crauste",
        plotted_complexity=5,
        data_path="Data/Data_fine/Crauste_fine.xlsx",
        model_path="Model_Crauste.xlsx",
        smooth_window=11,
        smooth_order=2,
        full_library_uses_smoothed_states=True,
        differential_order=1,
        thresholds=(1e-1, 1e-2, 1e-7, 1e-14),
    ),
    ModelConfig(
        name="Becker",
        plotted_complexity=6,
        data_path="Data/Data_fine/Becker_fine.xlsx",
        model_path="Model_Becker.xlsx",
        smooth_window=5,
        smooth_order=3,
        full_library_uses_smoothed_states=False,
        differential_order=2,
        thresholds=(1e-9, 1e-20, 1e-13, 1e-15),
        alpha=0.1,
    ),
    ModelConfig(
        name="Sneyd",
        plotted_complexity=6,
        data_path="Data/Data_fine/Sneyd_fine.xlsx",
        model_path="Model_Sneyd.xlsx",
        smooth_window=11,
        smooth_order=3,
        full_library_uses_smoothed_states=True,
        differential_order=2,
        # The notebook defines a degree-specific dictionary but its actual fit
        # call uses 1e-4 for all four degrees.
        thresholds=(1e-4, 1e-4, 1e-4, 1e-4),
    ),
    ModelConfig(
        name="Boehm",
        plotted_complexity=8,
        data_path="Data/Data_fine/Boehm_fine.xlsx",
        model_path="Model_Boehm.xlsx",
        smooth_window=7,
        smooth_order=2,
        full_library_uses_smoothed_states=True,
        differential_order=2,
        # As in the executed notebook call (not its unused threshold dict).
        thresholds=(1e-8, 1e-8, 1e-8, 1e-8),
    ),
    ModelConfig(
        name="Hass",
        # Plot_3 labels Hass as n=9.  Its actual analysis removes two inhibited
        # state columns and therefore constructs a seven-state library.  The
        # historical plot label is retained for direct visual comparison.
        plotted_complexity=9,
        data_path="Data_used/Hass.xlsx",
        data_source="shared",
        model_path="Model_Hass.xlsx",
        smooth_window=11,
        smooth_order=2,
        full_library_uses_smoothed_states=True,
        differential_order=2,
        thresholds=(1e-8, 1e-8, 1e-8, 1e-8),
        drop_state_columns=(
            "ApoER2_Dab1SFKInh",
            "ApoER2_re_Dab1SFKInh",
        ),
    ),
    ModelConfig(
        name="Swameye",
        plotted_complexity=9,
        data_path="Data/Data_fine/Swameye_fine.xlsx",
        model_path="Model_Swameye.xlsx",
        smooth_window=11,
        smooth_order=3,
        full_library_uses_smoothed_states=True,
        differential_order=2,
        thresholds=(1e-3, 1e-3, 1e-3, 1e-3),
        recover_from_unsmoothed_data=True,
    ),
    ModelConfig(
        name="Zheng",
        plotted_complexity=15,
        data_path="Data/Data_fine/Zheng_fine.xlsx",
        model_path="Model_Zheng.xlsx",
        smooth_window=7,
        smooth_order=3,
        full_library_uses_smoothed_states=True,
        differential_order=2,
        # The executed notebook fit hard-codes 1e-5 for every degree.
        thresholds=(1e-5, 1e-5, 1e-5, 1e-5),
    ),
)


def configs_for_dataset(
    configs: Sequence[ModelConfig], dataset: str
) -> tuple[ModelConfig, ...]:
    """Select real Exp/Fine files; label Hass's common source explicitly."""
    if dataset not in DATASETS:
        raise ValueError(f"Unknown dataset: {dataset!r}")
    return tuple(
        replace(
            config,
            dataset=dataset,
            data_source="shared" if config.name == "Hass" else dataset,
            data_path=(
                "Data_used/Hass.xlsx"
                if config.name == "Hass"
                else f"Data/Data_{dataset}/{config.name}_{dataset}.xlsx"
            ),
        )
        for config in configs
    )


@dataclass
class PreparedModel:
    config: ModelConfig
    original_data: pd.DataFrame
    raw_states: pd.DataFrame
    smoothed_data: pd.DataFrame
    smoothed_states: pd.DataFrame
    derivatives: pd.DataFrame
    original_model: pd.DataFrame


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Recompute Plot_3 Figures 1–4 separately for Exp and Fine data, "
            "using L2 library-column normalization by default."
        )
    )
    parser.add_argument(
        "--sindy-root",
        type=Path,
        default=DEFAULT_SINDY_ROOT,
        help=(
            "Directory containing models_set2 (default: inferred from this "
            "script's location)."
        ),
    )
    parser.add_argument(
        "--models",
        nargs="+",
        choices=[config.name for config in MODEL_CONFIGS],
        default=[config.name for config in MODEL_CONFIGS],
        help="Models to run; the default is all eight Plot_3 models.",
    )
    parser.add_argument(
        "--datasets",
        nargs="+",
        choices=DATASETS,
        default=list(DATASETS),
        help="Data sources to run separately (default: exp fine). Hass uses shared data.",
    )
    parser.add_argument(
        "--degrees",
        nargs="+",
        type=int,
        choices=DEGREES,
        default=list(DEGREES),
        help="Polynomial degrees to run; the default is 2 3 4 5.",
    )
    parser.add_argument(
        "--scaling",
        choices=("std", "l2"),
        default="l2",
        help=(
            "Library-column scaling used only for condition numbers: 'std' "
            "gives sample std=1, while 'l2' gives column L2 norm=1 (default). "
            "Sparse regression remains on the original library scale."
        ),
    )
    parser.add_argument(
        "--center",
        action="store_true",
        help=(
            "Also subtract each library-column mean before dividing by its "
            "sample standard deviation; valid only with --scaling std."
        ),
    )
    parser.add_argument(
        "--term-cache",
        type=Path,
        default=None,
        help=(
            "Optional JSON cache for wrong/missing term sets. Existing entries "
            "skip SINDy recovery; newly computed entries are written back."
        ),
    )
    parser.add_argument(
        "--full-only",
        action="store_true",
        help="Compute full-library tables and Figures 1/2 only; skip SINDy recovery.",
    )
    parser.add_argument(
        "--save",
        action="store_true",
        help="Save PNG/SVG figures and one Excel workbook per dataset, plus term cache.",
    )
    parser.add_argument(
        "--no-show",
        action="store_true",
        help="Do not open figure windows (normally used with --save).",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Output directory used by --save.",
    )
    args = parser.parse_args(argv)
    if args.no_show and not args.save:
        parser.error("--no-show would hide all outputs; also pass --save.")
    if args.center and args.scaling != "std":
        parser.error("--center is only valid with --scaling std.")
    args.degrees = tuple(sorted(set(args.degrees)))
    args.datasets = tuple(dict.fromkeys(args.datasets))
    return args


def configure_runtime(no_show: bool) -> None:
    """Set writable caches before importing repository plotting modules."""
    cache_root = Path(tempfile.gettempdir()) / "sindy_plot3_condition_scaling_cache"
    (cache_root / "matplotlib").mkdir(parents=True, exist_ok=True)
    (cache_root / "xdg").mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(cache_root / "matplotlib"))
    os.environ.setdefault("XDG_CACHE_HOME", str(cache_root / "xdg"))
    if no_show:
        os.environ.setdefault("MPLBACKEND", "Agg")

    if str(SCRIPT_DIR) not in sys.path:
        sys.path.insert(0, str(SCRIPT_DIR))

    warnings.filterwarnings(
        "ignore",
        message="pkg_resources is deprecated as an API.*",
        category=UserWarning,
    )


def load_repository_modules():
    import Comparison
    from Recover_Model import Recover_Model_sindy
    from dae_finder import PolyFeatureMatrix

    return Comparison, Recover_Model_sindy, PolyFeatureMatrix


def _as_time_columns(time_column: str | Sequence[str]) -> list[str]:
    return [time_column] if isinstance(time_column, str) else list(time_column)


def prepare_model(config: ModelConfig, model_root: Path, Comparison) -> PreparedModel:
    data_file = model_root / config.data_path
    model_file = model_root / config.model_path
    if not data_file.exists():
        raise FileNotFoundError(f"Missing {config.name} data file: {data_file}")
    if not model_file.exists():
        raise FileNotFoundError(f"Missing {config.name} model file: {model_file}")

    original = pd.read_excel(data_file)
    if config.drop_state_columns:
        missing = set(config.drop_state_columns) - set(original.columns)
        if missing:
            raise KeyError(f"{config.name}: missing columns requested for removal: {missing}")
        original = original.drop(columns=list(config.drop_state_columns))

    renamed, time_column = Comparison.standardize_columns(original)
    time_columns = _as_time_columns(time_column)
    if config.minimum_time is not None:
        time_name = time_columns[0]
        original = original.loc[original[time_name] >= config.minimum_time].reset_index(
            drop=True
        )
        renamed, time_column = Comparison.standardize_columns(original)
        time_columns = _as_time_columns(time_column)

    raw_states = renamed.drop(columns=time_columns).copy()
    expected_names = [f"x{index}" for index in range(1, raw_states.shape[1] + 1)]
    raw_states.columns = expected_names

    smoothed, derivatives_with_time = Comparison.sg_smooth_and_cdiff(
        original,
        window_length=config.smooth_window,
        polyorder=config.smooth_order,
    )
    smoothed_renamed, smoothed_time = Comparison.standardize_columns(smoothed)
    smoothed_states = smoothed_renamed.drop(
        columns=_as_time_columns(smoothed_time)
    ).copy()
    smoothed_states.columns = expected_names
    derivatives = derivatives_with_time.drop(columns=time_columns).copy()

    for label, frame in (
        ("raw states", raw_states),
        ("smoothed states", smoothed_states),
        ("derivatives", derivatives),
    ):
        values = frame.to_numpy(dtype=float)
        if values.ndim != 2 or not np.all(np.isfinite(values)):
            raise ValueError(f"{config.name}: {label} contains invalid values.")
    if raw_states.shape != smoothed_states.shape or raw_states.shape != derivatives.shape:
        raise ValueError(
            f"{config.name}: state/derivative shapes do not match: "
            f"raw={raw_states.shape}, smooth={smoothed_states.shape}, "
            f"derivative={derivatives.shape}."
        )

    return PreparedModel(
        config=config,
        original_data=original,
        raw_states=raw_states,
        smoothed_data=smoothed,
        smoothed_states=smoothed_states,
        derivatives=derivatives,
        original_model=pd.read_excel(model_file),
    )


def build_full_library(states: pd.DataFrame, degree: int, PolyFeatureMatrix) -> pd.DataFrame:
    library = PolyFeatureMatrix(degree).fit_transform(states)
    if "1" not in library.columns:
        raise KeyError("Polynomial library does not contain the expected constant column '1'.")
    library = library.drop(columns=["1"])
    expected = math.comb(states.shape[1] + degree, degree) - 1
    if library.shape != (len(states), expected):
        raise AssertionError(
            f"Expected library shape {(len(states), expected)}, got {library.shape}."
        )
    return library


def scaling_details(scaling: str, center: bool) -> tuple[str, str, str]:
    """Return the CSV label, figure label, and filename suffix."""
    if scaling == "l2":
        if center:
            raise ValueError("Centering is not part of L2 column scaling.")
        return "unit_l2", "library-column L2=1", "l2"
    if scaling == "std":
        if center:
            return "centered_zscore", "centered library columns, std=1", "zscore"
        return "std_only", "library-column std=1, no centering", "std1"
    raise ValueError(f"Unknown scaling method: {scaling!r}")


def scale_library_columns(
    values: np.ndarray,
    *,
    scaling: str,
    center: bool,
) -> np.ndarray:
    """Scale every library column independently using the selected method."""
    values = np.asarray(values, dtype=float)
    if values.ndim != 2 or values.shape[0] < 2 or values.shape[1] < 1:
        raise ValueError(f"Expected a nonempty 2D matrix with at least two rows; got {values.shape}.")
    scaling_details(scaling, center)  # validate the combination

    if scaling == "std":
        divisors = values.std(axis=0, ddof=1)
        numerator = values - values.mean(axis=0) if center else values
        diagnostic = "sample standard deviation"
    else:
        divisors = np.linalg.norm(values, ord=2, axis=0)
        numerator = values
        diagnostic = "L2 norm"

    invalid = (~np.isfinite(divisors)) | (divisors <= 0.0)
    if np.any(invalid):
        bad = np.flatnonzero(invalid).tolist()
        raise ValueError(
            f"Cannot scale columns with zero/non-finite {diagnostic} at "
            f"indices {bad[:10]}."
        )
    scaled = numerator / divisors
    if not np.all(np.isfinite(scaled)):
        raise ValueError(f"Unit-{diagnostic} scaling produced non-finite values.")

    if scaling == "std":
        check = scaled.std(axis=0, ddof=1)
    else:
        check = np.linalg.norm(scaled, ord=2, axis=0)
    if not np.allclose(check, 1.0, rtol=1e-10, atol=1e-12):
        raise AssertionError(f"A scaled library column does not have unit {diagnostic}.")
    return scaled


def condition_diagnostics(values: np.ndarray) -> tuple[float, int]:
    """Return the 2-norm condition number and numerical matrix rank.

    Computing both from the same singular values avoids a second expensive SVD,
    which matters for the 15-state, degree-5 Zheng library.
    Like numpy.linalg.cond, this uses min(n_rows, n_columns) singular values.
    A finite ratio for a wide library does not imply full column rank.
    """
    matrix = np.asarray(values, dtype=float)
    singular_values = np.linalg.svd(matrix, compute_uv=False)
    condition = float(singular_values[0] / singular_values[-1])
    if math.isnan(condition) or condition < 1.0:
        raise FloatingPointError(f"Invalid 2-norm condition number: {condition}")
    tolerance = (
        singular_values[0]
        * max(matrix.shape)
        * np.finfo(singular_values.dtype).eps
    )
    numerical_rank = int(np.count_nonzero(singular_values > tolerance))
    return condition, numerical_rank


def _term_exponent(term: str, n_states: int) -> tuple[int, ...]:
    text = str(term).strip().replace("*", " ")
    if text in ("", "1"):
        return (0,) * n_states
    exponent = [0] * n_states
    for token in text.split():
        match = re.fullmatch(r"x(\d+)(?:\^(\d+))?", token)
        if match is None:
            raise ValueError(f"Unsupported polynomial term format: {term!r}")
        index = int(match.group(1)) - 1
        if index < 0 or index >= n_states:
            raise ValueError(f"Term {term!r} refers to unavailable state x{index + 1}.")
        exponent[index] += int(match.group(2) or 1)
    return tuple(exponent)


def evaluate_terms(states: pd.DataFrame, terms: Sequence[str]) -> tuple[np.ndarray, list[str]]:
    """Evaluate selected monomials without constructing the entire library."""
    n_states = states.shape[1]
    unique: dict[tuple[int, ...], str] = {}
    for term in terms:
        exponent = _term_exponent(term, n_states)
        if not any(exponent):
            # The Plot_3 candidate libraries drop the constant column.
            continue
        unique.setdefault(exponent, str(term))
    ordered = sorted(unique.items())
    if not ordered:
        return np.empty((len(states), 0), dtype=float), []

    state_values = states.to_numpy(dtype=float)
    columns = []
    labels = []
    for exponent, label in ordered:
        column = np.ones(len(states), dtype=float)
        for state_index, power in enumerate(exponent):
            if power:
                column *= state_values[:, state_index] ** power
        columns.append(column)
        labels.append(label)
    return np.column_stack(columns), labels


class _TermsOnlyMixin:
    """Suppress Terms_Analysis' old internal SVD; only term sets are needed."""

    def _SVD_analysis(self):
        return {}


def recover_mixed_terms(
    prepared: PreparedModel,
    degree: int,
    Comparison,
    Recover_Model_sindy,
) -> dict[str, list[str]]:
    """Recover wrong/missing terms using the original, unscaled library."""
    config = prepared.config
    fit_data = (
        prepared.original_data
        if config.recover_from_unsmoothed_data
        else prepared.smoothed_data
    )
    xdot = None if config.recover_from_unsmoothed_data else prepared.derivatives

    # The wrapper reports feature names on stdout. Keep this standalone script's
    # progress output concise while allowing warnings/errors to remain visible.
    with contextlib.redirect_stdout(io.StringIO()):
        recovered = Recover_Model_sindy(
            fit_data,
            differential_order=config.differential_order,
            poly_degree=degree,
            threshold=config.threshold(degree),
            alpha=config.alpha,
            after_threshold=config.after_threshold,
            Xdot=xdot,
            method="Monomial",
            L=prepared.raw_states.min(axis=0),
            U=prepared.raw_states.max(axis=0),
            denormalize=False,
        )

    identifier = Comparison.Terms_Identification_original(prepared.original_model)
    variable_mapping, _ = identifier._parse_model()
    TermsOnlyAnalysis = type(
        "TermsOnlyAnalysis",
        (_TermsOnlyMixin, Comparison.Terms_Analysis),
        {},
    )
    analyzer = TermsOnlyAnalysis(
        prepared.original_model,
        recovered.model_expression,
        variable_mapping,
        pd.DataFrame(index=prepared.raw_states.index),
        degree,
    )
    result = {
        equation: sorted(str(term) for term in terms)
        for equation, terms in analyzer.Mixed_terms.items()
    }
    del analyzer, recovered
    gc.collect()
    return result


def load_term_cache(path: Path | None) -> dict[str, Any]:
    if path is None or not path.exists():
        return {}
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError(f"Term cache must contain a JSON object: {path}")
    if data and not any(dataset in data for dataset in DATASETS):
        # The previous script only analyzed Fine data. Never reuse these
        # recovered supports for Exp, even when model and degree are identical.
        model_names = {config.name for config in MODEL_CONFIGS}
        if not set(data).issubset(model_names):
            raise ValueError(f"Unrecognized legacy term-cache keys: {path}")
        data = {"fine": data}
    if not set(data).issubset(DATASETS) or any(
        not isinstance(entries, dict) for entries in data.values()
    ):
        raise ValueError(f"Expected dataset -> model -> degree term cache: {path}")
    return data


def write_term_cache(path: Path, cache: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(cache, handle, indent=2, sort_keys=True)


def analyze(
    configs: Sequence[ModelConfig],
    degrees: Sequence[int],
    model_root: Path,
    scaling: str,
    center: bool,
    term_cache: dict[str, Any],
    Comparison,
    Recover_Model_sindy,
    PolyFeatureMatrix,
    *,
    full_only: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    full_rows: list[dict[str, Any]] = []
    mixed_rows: list[dict[str, Any]] = []
    scaling_name, _, _ = scaling_details(scaling, center)

    for config in configs:
        model_start = time.perf_counter()
        prepared = prepare_model(config, model_root, Comparison)
        actual_states = prepared.raw_states.shape[1]
        print(
            f"[{config.dataset}/{config.name}] {len(prepared.raw_states)} rows, "
            f"{actual_states} analyzed states (plot label n={config.plotted_complexity})",
            flush=True,
        )
        recovery_cache = (
            {} if full_only else term_cache.setdefault(config.name, {})
        )
        if not isinstance(recovery_cache, dict):
            raise ValueError(f"Invalid term cache for {config.name}: expected an object.")

        for degree in degrees:
            degree_start = time.perf_counter()
            full_states = (
                prepared.smoothed_states
                if config.full_library_uses_smoothed_states
                else prepared.raw_states
            )
            print(f"  degree {degree}: full library", end="", flush=True)
            full_library = build_full_library(full_states, degree, PolyFeatureMatrix)
            # Figures 1/2: build the complete library first, then scale only the
            # matrix whose condition number is being evaluated.
            scaled_full = scale_library_columns(
                full_library.to_numpy(dtype=float),
                scaling=scaling,
                center=center,
            )
            full_condition, full_rank = condition_diagnostics(scaled_full)
            full_rows.append(
                {
                    "dataset": config.dataset,
                    "data_source": config.data_source,
                    "data_file": str((model_root / config.data_path).resolve()),
                    "model": config.name,
                    "plotted_complexity": config.plotted_complexity,
                    "actual_n_states": actual_states,
                    "degree": degree,
                    "n_observations": len(full_library),
                    "n_library_terms": full_library.shape[1],
                    "is_wide_library": full_library.shape[1] > len(full_library),
                    "numerical_rank": full_rank,
                    "full_column_rank": full_rank == full_library.shape[1],
                    "state_preprocessing": (
                        "smoothed" if config.full_library_uses_smoothed_states else "raw"
                    ),
                    "scaling": scaling_name,
                    "condition_number": full_condition,
                    "log10_condition_number": math.log10(full_condition),
                }
            )
            del full_library, scaled_full
            gc.collect()

            if full_only:
                print(f"; done in {time.perf_counter() - degree_start:.1f}s", flush=True)
                continue

            degree_key = str(degree)
            if degree_key in recovery_cache:
                mixed_terms = recovery_cache[degree_key]
                print(
                    "; cached wrong/missing terms",
                    end="",
                    flush=True,
                )
            else:
                print(
                    "; recovering wrong/missing terms on original library scale",
                    end="",
                    flush=True,
                )
                # Figures 3/4: SINDy and support comparison use the original
                # feature scale. Only the fixed sublibrary is scaled below.
                mixed_terms = recover_mixed_terms(
                    prepared,
                    degree,
                    Comparison,
                    Recover_Model_sindy,
                )
                recovery_cache[degree_key] = mixed_terms

            for equation, terms in mixed_terms.items():
                sublibrary, evaluated_terms = evaluate_terms(prepared.raw_states, terms)
                if sublibrary.shape[1] == 0:
                    # Match Plot_3/Table_fine's convention: an exactly recovered
                    # equation is represented by log10(condition)=0.
                    sub_condition = 1.0
                    log_condition = 0.0
                    sub_rank = 0
                else:
                    # Scale the already identified wrong/missing-term
                    # sublibrary immediately before computing its condition.
                    scaled_sub = scale_library_columns(
                        sublibrary,
                        scaling=scaling,
                        center=center,
                    )
                    sub_condition, sub_rank = condition_diagnostics(scaled_sub)
                    log_condition = math.log10(sub_condition)
                mixed_rows.append(
                    {
                        "dataset": config.dataset,
                        "data_source": config.data_source,
                        "data_file": str((model_root / config.data_path).resolve()),
                        "model": config.name,
                        "plotted_complexity": config.plotted_complexity,
                        "actual_n_states": actual_states,
                        "degree": degree,
                        "equation": equation,
                        "n_wrong_or_missing_terms": len(evaluated_terms),
                        "numerical_rank": sub_rank,
                        "wrong_or_missing_terms": " | ".join(evaluated_terms),
                        "sparse_regression_scaling": "original_scale",
                        "scaling": scaling_name,
                        "condition_number": sub_condition,
                        "log10_condition_number": log_condition,
                    }
                )

            elapsed = time.perf_counter() - degree_start
            print(f"; done in {elapsed:.1f}s", flush=True)

        print(
            f"[{config.name}] completed in {(time.perf_counter() - model_start) / 60:.1f} min",
            flush=True,
        )

    return pd.DataFrame(full_rows), pd.DataFrame(mixed_rows), term_cache


def configure_plot_style() -> None:
    import matplotlib as mpl
    import matplotlib.pyplot as plt

    plt.style.use("seaborn-v0_8-whitegrid")
    mpl.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["DejaVu Sans", "Arial", "Helvetica"],
            "font.size": 20,
            "axes.linewidth": 1.0,
            "grid.linewidth": 0.5,
            "xtick.major.width": 1.0,
            "ytick.major.width": 1.0,
            "legend.frameon": True,
            "legend.shadow": False,
            "legend.edgecolor": "black",
            "legend.fancybox": False,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
        }
    )


def _ordered_models(configs: Sequence[ModelConfig]) -> list[ModelConfig]:
    order = {config.name: index for index, config in enumerate(MODEL_CONFIGS)}
    return sorted(configs, key=lambda config: (config.plotted_complexity, order[config.name]))


def _style_axis(axis, title: str) -> None:
    from matplotlib.ticker import MultipleLocator

    axis.set_xlabel("Model (sorted by complexity)")
    axis.set_ylabel(r"$\log_{10}$(Cond Number)")
    axis.set_title(title, pad=15)
    axis.grid(True, alpha=0.4, linestyle="--", linewidth=0.6, color="gray")
    axis.yaxis.set_minor_locator(MultipleLocator(5))
    axis.tick_params(which="both", direction="in", top=True, right=True)
    axis.tick_params(which="major", length=6, width=1.0)
    axis.tick_params(which="minor", length=3, width=0.5)


def plot_full_condition(
    full_results: pd.DataFrame,
    configs: Sequence[ModelConfig],
    degrees: Sequence[int],
    scaling: str,
    center: bool,
):
    import matplotlib.pyplot as plt

    ordered = _ordered_models(configs)
    model_names = [config.name for config in ordered]
    x = np.arange(len(ordered))
    fig, axis = plt.subplots(figsize=(10, 8))
    for degree_index, degree in enumerate(degrees):
        indexed = full_results[full_results["degree"] == degree].set_index("model")
        y = indexed.loc[model_names, "log10_condition_number"].to_numpy(float)
        style_index = DEGREES.index(degree)
        axis.plot(
            x,
            y,
            marker=DEGREE_MARKERS[style_index],
            color=DEGREE_COLORS[style_index],
            label=f"Degree = {degree}",
            linewidth=1.8,
            linestyle=DEGREE_LINESTYLES[style_index],
            markersize=7,
            markeredgewidth=1.0,
            markeredgecolor="white",
            alpha=0.85,
        )
    axis.set_xticks(x)
    axis.set_xticklabels(
        [f"{config.name}\n({config.plotted_complexity})" for config in ordered],
        fontsize=15,
    )
    _, variant, _ = scaling_details(scaling, center)
    _style_axis(
        axis,
        f"Condition Number of Full Library",
        #f"Full Library Condition Number vs Model Complexity\n({variant})",
    )
    legend = axis.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, 1.20),
        ncol=4,
        frameon=True,
        framealpha=0.95,
        fontsize=15,
        handlelength=2.5,
        borderpad=0.5,
        edgecolor="black",
        facecolor="white",
    )
    legend.get_frame().set_linewidth(0.8)
    fig.tight_layout()
    return fig


def summarize_mixed_conditions(mixed_results: pd.DataFrame) -> pd.DataFrame:
    """Summarize equation-wise conditions for each dataset/model/degree.

    The arithmetic mean of condition numbers and the mean of their logarithms
    are distinct statistics. Keep both, with min/max on each scale. As in the
    existing plots, equations with no wrong/missing terms contribute condition
    1 (log10=0); they are not excluded from these summaries.
    """
    return (
        mixed_results.groupby(
            [
                "dataset",
                "data_source",
                "data_file",
                "model",
                "plotted_complexity",
                "actual_n_states",
                "degree",
                "scaling",
            ],
            as_index=False,
        )
        .agg(
            n_equations=("equation", "nunique"),
            mean_condition_number=("condition_number", "mean"),
            min_condition_number=("condition_number", "min"),
            max_condition_number=("condition_number", "max"),
            mean_log10_condition=("log10_condition_number", "mean"),
            min_log10_condition=("log10_condition_number", "min"),
            max_log10_condition=("log10_condition_number", "max"),
            std_log10_condition=(
                "log10_condition_number",
                lambda values: values.std(ddof=1) if len(values) > 1 else 0.0,
            ),
        )
    )


def plot_mixed_condition(
    mixed_summary: pd.DataFrame,
    configs: Sequence[ModelConfig],
    degrees: Sequence[int],
    scaling: str,
    center: bool,
):
    import matplotlib.pyplot as plt

    ordered = _ordered_models(configs)
    model_names = [config.name for config in ordered]
    x = np.arange(len(ordered))
    fig, axis = plt.subplots(figsize=(10, 8))
    for degree in degrees:
        indexed = mixed_summary[mixed_summary["degree"] == degree].set_index("model")
        y = indexed.loc[model_names, "mean_log10_condition"].to_numpy(float)
        yerr = indexed.loc[model_names, "std_log10_condition"].to_numpy(float)
        style_index = DEGREES.index(degree)
        axis.errorbar(
            x,
            y,
            yerr=yerr,
            fmt=DEGREE_MARKERS[style_index] + DEGREE_LINESTYLES[style_index],
            color=DEGREE_COLORS[style_index],
            label=f"Degree = {degree}",
            linewidth=1.8,
            markersize=6,
            capsize=3,
            elinewidth=1.0,
            capthick=1.0,
        )
    axis.set_xticks(x)
    axis.set_xticklabels(
        [f"{config.name}\n({config.plotted_complexity})" for config in ordered],
        fontsize=15,
    )
    _, variant, _ = scaling_details(scaling, center)
    _style_axis(
        axis,
        f"Condition Number of Incorrect and Missing Terms",
        #f"Wrong & Missing Terms Condition Number vs Model Complexity\n({variant})",
    )
    legend = axis.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, 1.20),
        ncol=4,
        frameon=True,
        framealpha=0.95,
        fontsize=15,
        handlelength=2.5,
        borderpad=0.5,
        edgecolor="black",
        facecolor="white",
    )
    legend.get_frame().set_linewidth(0.8)
    fig.tight_layout()
    return fig


def plot_full_condition_by_degree(
    full_results: pd.DataFrame,
    configs: Sequence[ModelConfig],
    degrees: Sequence[int],
    scaling: str,
    center: bool,
):
    """Plot Figure 1 from the same scaled libraries used in Figure 2."""
    import matplotlib.pyplot as plt

    model_colors = (
        "#004C6D", "#A86500", "#006644", "#9B4D8C",
        "#B8A800", "#0066CC", "#AA4400", "#AA4400",
    )
    model_markers = ("s", "^", "v", "D", "p", "*", "h", "X")
    style_indices = {config.name: index for index, config in enumerate(MODEL_CONFIGS)}
    degree_order = list(degrees)
    fig, axis = plt.subplots(figsize=(10, 8))
    for config in _ordered_models(configs):
        indexed = full_results[full_results["model"] == config.name].set_index("degree")
        y = indexed.loc[degree_order, "log10_condition_number"].to_numpy(float)
        style_index = style_indices[config.name]
        axis.plot(
            degree_order,
            y,
            marker=model_markers[style_index],
            color=model_colors[style_index],
            label=f"{config.name} (n={config.plotted_complexity})",
            linewidth=1.8,
            markersize=7,
            markeredgewidth=1.0,
            markeredgecolor="white",
            alpha=0.85,
        )
    _style_axis(axis, "Full Library Condition Number vs. Library Degree")
    axis.set_xlabel("Degree")
    axis.set_xticks(degree_order)
    axis.set_xlim(min(degree_order) - 0.2, max(degree_order) + 0.2)
    legend = axis.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, 1.12),
        ncol=2,
        frameon=True,
        framealpha=0.95,
        fontsize=15,
        columnspacing=1.0,
        handlelength=1.5,
        borderpad=0.5,
        edgecolor="black",
        facecolor="white",
    )
    legend.get_frame().set_linewidth(0.8)
    fig.tight_layout()
    return fig


def plot_mixed_condition_by_degree(
    mixed_summary: pd.DataFrame,
    configs: Sequence[ModelConfig],
    degrees: Sequence[int],
    scaling: str,
    center: bool,
):
    """Plot Figure 3: equation-wise log10 mean and sample SD per degree."""
    import matplotlib.pyplot as plt
    from matplotlib.ticker import MultipleLocator

    model_colors = (
        "#004C6D", "#A86500", "#006644", "#9B4D8C",
        "#B8A800", "#0066CC", "#AA4400", "#AA4400",
    )
    model_markers = ("s", "^", "v", "D", "p", "*", "h", "X")
    style_indices = {config.name: index for index, config in enumerate(MODEL_CONFIGS)}
    degree_order = list(degrees)
    fig, axis = plt.subplots(figsize=(10, 8))
    for config in _ordered_models(configs):
        indexed = mixed_summary[mixed_summary["model"] == config.name].set_index("degree")
        y = indexed.loc[degree_order, "mean_log10_condition"].to_numpy(float)
        yerr = indexed.loc[degree_order, "std_log10_condition"].to_numpy(float)
        style_index = style_indices[config.name]
        axis.errorbar(
            degree_order,
            y,
            yerr=yerr,
            fmt=model_markers[style_index] + "-",
            color=model_colors[style_index],
            label=f"{config.name} (n={config.plotted_complexity})",
            linewidth=1.8,
            markersize=7,
            markeredgewidth=1.0,
            markeredgecolor="white",
            capsize=3,
            elinewidth=1.0,
            capthick=1.0,
            alpha=0.85,
        )
    _style_axis(axis, "Wrong & Missing Terms Condition Number vs Library Degree")
    axis.set_xlabel("Degree")
    axis.set_xticks(degree_order)
    axis.set_xlim(min(degree_order) - 0.2, max(degree_order) + 0.2)
    axis.yaxis.set_minor_locator(MultipleLocator(2.5))
    legend = axis.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, 1.12),
        ncol=2,
        frameon=True,
        framealpha=0.95,
        fontsize=15,
        columnspacing=1.0,
        handlelength=1.5,
        borderpad=0.5,
        edgecolor="black",
        facecolor="white",
    )
    legend.get_frame().set_linewidth(0.8)
    fig.tight_layout()
    return fig


def plot_results(
    full_results: pd.DataFrame,
    mixed_summary: pd.DataFrame,
    configs: Sequence[ModelConfig],
    degrees: Sequence[int],
    scaling: str,
    center: bool,
) -> dict[str, Any]:
    """Create Figures 1–4 from one dataset's already-computed results."""
    dataset = str(full_results["dataset"].iloc[0])
    _, scaling_label, _ = scaling_details(scaling, center)
    plot_args = (configs, degrees, scaling, center)
    figures = {
        "Cond_full_degree": plot_full_condition_by_degree(full_results, *plot_args),
        "Cond_full_model": plot_full_condition(full_results, *plot_args),
    }
    if not mixed_summary.empty:
        figures.update({
            "wrong_missing_degree": plot_mixed_condition_by_degree(
                mixed_summary, *plot_args
            ),
            "wrong_missing_model": plot_mixed_condition(mixed_summary, *plot_args),
        })
    for name, figure in figures.items():
        axis = figure.axes[0]
        axis.set_title(
            f"{axis.get_title()}\n{dataset.capitalize()} data · {scaling_label}",
            fontsize=18,
            pad=15,
        )
        figure.canvas.manager.set_window_title(f"Plot_3 | {dataset} | {name}")
        if name.endswith("degree"):
            axis.get_legend().set_bbox_to_anchor((0.5, 1.26))
        figure.tight_layout()
    return figures


def save_excel_results(
    path: Path,
    full_results: pd.DataFrame,
    mixed_results: pd.DataFrame,
    mixed_summary: pd.DataFrame,
) -> None:
    """Export numeric results without rounding; formatting affects display only."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    model_order = full_results["model"].drop_duplicates()
    sheets: dict[str, pd.DataFrame] = {}
    for name, value in (
        ("Full_condition", "condition_number"),
        ("Full_log10", "log10_condition_number"),
    ):
        table = full_results.pivot(index="model", columns="degree", values=value)
        table = table.reindex(model_order).sort_index(axis=1)
        table.columns = [f"Degree {degree}" for degree in table.columns]
        sheets[name] = table.rename_axis("Model").reset_index()
    sheets["Full_details"] = full_results
    if not mixed_results.empty:
        sheets["Wrong_missing_summary"] = mixed_summary
        sheets["Wrong_missing_by_equation"] = mixed_results

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False, inf_rep="inf")
            sheet = writer.sheets[name]
            sheet.sheet_view.showGridLines = False
            sheet.auto_filter.ref = sheet.dimensions
            sheet.freeze_panes = "B2"
            sheet.row_dimensions[1].height = 34
            for column_index, column in enumerate(frame.columns, start=1):
                label = str(column)
                letter = get_column_letter(column_index)
                width = max(18, min(32, len(label) + 2))
                if label in ("data_file", "wrong_or_missing_terms"):
                    width = 80
                sheet.column_dimensions[letter].width = width
                header = sheet.cell(1, column_index)
                header.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                header.fill = PatternFill("solid", fgColor="24476B")
                header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                is_condition = label in {
                    "condition_number", "mean_condition_number",
                    "min_condition_number", "max_condition_number",
                } or (
                    name == "Full_condition" and column_index > 1
                )
                is_log = "log10" in label or (name == "Full_log10" and column_index > 1)
                for cells in sheet.iter_rows(
                    min_row=2, min_col=column_index, max_col=column_index
                ):
                    cell = cells[0]
                    cell.font = Font(name="Arial", size=11)
                    cell.alignment = Alignment(vertical="center")
                    if is_condition:
                        cell.number_format = "0.000000E+00"
                    elif is_log:
                        cell.number_format = "0.000000"
            for row_index in range(2, sheet.max_row + 1):
                sheet.row_dimensions[row_index].height = 20


def save_outputs(
    output_dir: Path,
    full_results: pd.DataFrame,
    mixed_results: pd.DataFrame,
    mixed_summary: pd.DataFrame,
    figures: dict[str, Any],
    scaling: str,
    center: bool,
) -> None:
    """Save one dataset in its own directory, including all model/degree rows."""
    output_dir.mkdir(parents=True, exist_ok=True)
    _, _, suffix = scaling_details(scaling, center)
    save_excel_results(
        output_dir / f"plot3_condition_numbers_{suffix}.xlsx",
        full_results, mixed_results, mixed_summary,
    )
    for name, figure in figures.items():
        export_options = {"bbox_inches": "tight", "pad_inches": 0.02}
        figure.savefig(output_dir / f"{name}_{suffix}.svg", **export_options)
        figure.savefig(output_dir / f"{name}_{suffix}.png", dpi=300, **export_options)
    print(f"Saved outputs to: {output_dir.resolve()}")


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    configure_runtime(args.no_show)
    Comparison, Recover_Model_sindy, PolyFeatureMatrix = load_repository_modules()

    model_root = args.sindy_root.expanduser().resolve() / "models_set2"
    if not model_root.is_dir():
        raise FileNotFoundError(
            f"Could not find models_set2 under --sindy-root: {model_root}"
        )
    selected_names = set(args.models)
    configs = tuple(config for config in MODEL_CONFIGS if config.name in selected_names)
    if args.term_cache is not None:
        cache_path = args.term_cache.expanduser().resolve()
    elif args.save:
        # Re-running the same --save command automatically reuses the expensive
        # SINDy-derived term sets produced by the first run.
        cache_path = (
            args.output_dir.expanduser().resolve()
            / "plot3_wrong_missing_terms.json"
        )
    else:
        cache_path = None
    term_cache = load_term_cache(cache_path)

    scaling_name, scaling_label, _ = scaling_details(args.scaling, args.center)
    if args.scaling == "l2":
        scaling_formula = "Theta[:, j] / ||Theta[:, j]||_2"
    elif args.center:
        scaling_formula = "(Theta[:, j] - mean_j) / sample_std_j"
    else:
        scaling_formula = "Theta[:, j] / sample_std_j"

    print("Plot_3 library-column-scaled condition-number experiment")
    print(f"  models: {', '.join(config.name for config in configs)}")
    print(f"  degrees: {args.degrees}")
    print(f"  datasets: {args.datasets}")
    if any(config.name == "Hass" for config in configs):
        print("  Hass: shared Data_used/Hass.xlsx source for Exp and Fine")
    print(f"  scaling: {scaling_name} ({scaling_label})")
    print(f"  formula: {scaling_formula}")
    print("  SINDy recovery: original, unscaled library columns")
    print(f"  data root: {model_root}")

    configure_plot_style()
    for dataset in args.datasets:
        dataset_configs = configs_for_dataset(configs, dataset)
        dataset_cache = term_cache.setdefault(dataset, {})
        full_results, mixed_results, dataset_cache = analyze(
            dataset_configs,
            args.degrees,
            model_root,
            args.scaling,
            args.center,
            dataset_cache,
            Comparison,
            Recover_Model_sindy,
            PolyFeatureMatrix,
            full_only=args.full_only,
        )
        term_cache[dataset] = dataset_cache
        mixed_summary = (
            summarize_mixed_conditions(mixed_results)
            if not mixed_results.empty else pd.DataFrame()
        )
        figures = plot_results(
            full_results, mixed_summary, dataset_configs,
            args.degrees, args.scaling, args.center,
        )

        print(f"\n{dataset.upper()} full-library log10 condition numbers:")
        print(
            full_results.pivot(
                index="model", columns="degree", values="log10_condition_number"
            )
            .reindex([config.name for config in _ordered_models(configs)])
            .round(4)
            .to_string()
        )
        if not mixed_summary.empty:
            print(f"\n{dataset.upper()} wrong/missing mean log10 condition numbers:")
            print(
                mixed_summary.pivot(
                    index="model", columns="degree", values="mean_log10_condition"
                )
                .reindex([config.name for config in _ordered_models(configs)])
                .round(4)
                .to_string()
            )
        if cache_path is not None and not args.full_only:
            write_term_cache(cache_path, term_cache)
            print(f"Updated dataset-specific term cache: {cache_path}")
        if args.save:
            save_outputs(
                args.output_dir.expanduser().resolve() / dataset,
                full_results,
                mixed_results,
                mixed_summary,
                figures,
                args.scaling,
                args.center,
            )
        if args.no_show:
            import matplotlib.pyplot as plt

            for figure in figures.values():
                plt.close(figure)
    if not args.no_show:
        import matplotlib.pyplot as plt

        plt.show()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
